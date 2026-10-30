"""
Live Excel A/B test sample-size calculator, seeded with this project's real
baseline conversion rate. Formula-driven (two-proportion z-test sample size
formula), not hardcoded — change the baseline or MDE and it recalculates.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.comments import Comment

wb = Workbook()
ws = wb.active
ws.title = "AB Test Calculator"

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
BOLD = Font(name=FONT, bold=True)
TITLE_FONT = Font(name=FONT, bold=True, size=14)
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1E2530")
YELLOW = PatternFill("solid", fgColor="FFFF00")
SECTION_FILL = PatternFill("solid", fgColor="EFEAE0")
thin = Side(style="thin", color="DDD6C7")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

ws.column_dimensions["A"].width = 42
for col in "BCDE":
    ws.column_dimensions[col].width = 16

def section(row, text, span=5):
    ws.cell(row=row, column=1, value=text).font = BOLD
    for col in range(1, span+1):
        ws.cell(row=row, column=col).fill = SECTION_FILL
    return row + 1

ws.merge_cells("A1:E1")
ws["A1"] = "Case 06 — A/B Test Sample-Size Calculator"
ws["A1"].font = TITLE_FONT
ws.merge_cells("A2:E2")
ws["A2"] = "Seeded with this project's real baseline (telephone-channel conversion rate). Edit any blue/yellow cell — every output recalculates."
ws["A2"].font = Font(name=FONT, italic=True, size=9, color="52514E")

r = 4
r = section(r, "1. BASELINE  (blue = from the observational analysis)")
row_baseline = r
ws.cell(row=r, column=1, value="Baseline conversion rate (lower-converting channel)").font = BLACK
c = ws.cell(row=r, column=2, value=0.0523); c.font = BLUE; c.number_format = "0.00%"; c.border = BORDER
c.comment = Comment("Telephone-channel conversion rate observed in the historical data — used as the conservative baseline for planning a future randomized test.", "BA Portfolio Playbook")
r += 1

r = section(r, "2. TEST DESIGN INPUTS  (yellow = edit for your own test)")
row_mde = r
ws.cell(row=r, column=1, value="Minimum detectable lift (percentage points)").font = BLACK
c = ws.cell(row=r, column=2, value=0.02); c.font = BLUE; c.fill = YELLOW; c.border = BORDER; c.number_format = "0.00%"
c.comment = Comment("The smallest true lift you want the test to reliably catch. Smaller = needs a much bigger sample (see the scenario table below).", "BA Portfolio Playbook")
row_z_alpha = r + 1
ws.cell(row=row_z_alpha, column=1, value="Z(alpha/2) for 95% significance, two-sided").font = BLACK
c = ws.cell(row=row_z_alpha, column=2, value=1.96); c.font = BLUE; c.fill = YELLOW; c.border = BORDER; c.number_format = "0.00"
c.comment = Comment("Standard values: 1.64 for 90% significance, 1.96 for 95%, 2.58 for 99%. Swap in your own if your test uses a different alpha.", "BA Portfolio Playbook")
row_z_beta = row_z_alpha + 1
ws.cell(row=row_z_beta, column=1, value="Z(beta) for 80% statistical power").font = BLACK
c = ws.cell(row=row_z_beta, column=2, value=0.84); c.font = BLUE; c.fill = YELLOW; c.border = BORDER; c.number_format = "0.00"
c.comment = Comment("Standard values: 0.84 for 80% power, 1.04 for 85%, 1.28 for 90%. Higher power needs a bigger sample.", "BA Portfolio Playbook")
r = row_z_beta + 2

r = section(r, "3. REQUIRED SAMPLE SIZE PER ARM  (formula — two-proportion z-test)")
row_p2 = r
ws.cell(row=r, column=1, value="Target conversion rate (baseline + MDE)").font = BLACK
c = ws.cell(row=r, column=2, value=f"=B{row_baseline}+B{row_mde}")
c.font = BLACK; c.number_format = "0.00%"; c.border = BORDER
row_n = r + 1
ws.cell(row=row_n, column=1, value="Required sample size per arm").font = BOLD
formula_n = (
    f"=((B{row_z_alpha}+B{row_z_beta})^2 * "
    f"(B{row_baseline}*(1-B{row_baseline}) + B{row_p2}*(1-B{row_p2}))) "
    f"/ (B{row_p2}-B{row_baseline})^2"
)
c = ws.cell(row=row_n, column=2, value=formula_n)
c.font = BOLD; c.number_format = "#,##0"; c.border = BORDER
row_total = row_n + 1
ws.cell(row=row_total, column=1, value="Total contacts needed (both arms)").font = BOLD
c = ws.cell(row=row_total, column=2, value=f"=B{row_n}*2")
c.font = BOLD; c.number_format = "#,##0"; c.border = BORDER
r = row_total + 2

r = section(r, "4. SCENARIO TABLE  (same formula, 3 MDE scenarios — for reference)")
scen_header = r
for col, label in enumerate(["MDE (pp)", "Target rate", "n per arm", "Total contacts"], start=1):
    ws.cell(row=r, column=col, value=label).font = HEADER_FONT
    ws.cell(row=r, column=col).fill = HEADER_FILL
r += 1
for mde_val in [0.01, 0.02, 0.03]:
    ws.cell(row=r, column=1, value=mde_val).number_format = "0.00%"; ws.cell(row=r, column=1).font = BLACK
    target_formula = f"=B{row_baseline}+A{r}"
    ws.cell(row=r, column=2, value=target_formula).number_format = "0.00%"; ws.cell(row=r, column=2).font = BLACK
    n_formula = (
        f"=((B{row_z_alpha}+B{row_z_beta})^2 * "
        f"(B{row_baseline}*(1-B{row_baseline}) + B{r}*(1-B{r}))) "
        f"/ (B{r}-B{row_baseline})^2"
    )
    c = ws.cell(row=r, column=3, value=n_formula); c.number_format = "#,##0"; c.font = BLACK; c.border = BORDER
    c2 = ws.cell(row=r, column=4, value=f"=C{r}*2"); c2.number_format = "#,##0"; c2.font = BLACK; c2.border = BORDER
    r += 1
r += 1

r = section(r, "LEGEND")
ws.cell(row=r, column=1, value="Blue text").font = BLACK
ws.cell(row=r, column=2, value="From the analysis, or an editable input").font = BLACK
r += 1
ws.cell(row=r, column=1, value="Yellow fill").font = BLACK
ws.cell(row=r, column=2, value="Key assumption — edit for your own test design").font = BLACK
r += 1
ws.cell(row=r, column=1, value="Black text").font = BLACK
ws.cell(row=r, column=2, value="Formula — recalculates automatically").font = BLACK

wb.save("case6_ab_test_calculator.xlsx")
print("Workbook saved.")
