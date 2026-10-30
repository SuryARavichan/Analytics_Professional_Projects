"""
Build the Case 02 cost-benefit / ROI model as a real Excel workbook with
live formulas (not hardcoded results), following the xlsx skill's
financial-model conventions: blue = input, black = formula, yellow = key
editable assumption.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "ROI Model"

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
BOLD = Font(name=FONT, bold=True)
TITLE_FONT = Font(name=FONT, bold=True, size=14)
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1E2530")
YELLOW = PatternFill("solid", fgColor="FFFF00")
SECTION_FILL = PatternFill("solid", fgColor="EFEAE0")
thin = Side(style="thin", color="DDD6C7")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

ws.column_dimensions["A"].width = 46
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 46

def section(row, text):
    ws.cell(row=row, column=1, value=text).font = BOLD
    for col in (1, 2):
        ws.cell(row=row, column=col).fill = SECTION_FILL
    return row + 1

def label_val(row, label, value, font=BLACK, fmt=None, comment=None, fill=None):
    lc = ws.cell(row=row, column=1, value=label)
    lc.font = BLACK
    vc = ws.cell(row=row, column=2, value=value)
    vc.font = font
    vc.border = BORDER
    if fmt:
        vc.number_format = fmt
    if fill:
        vc.fill = fill
    if comment:
        vc.comment = Comment(comment, "BA Portfolio Playbook")
    return row + 1

# ---- Title ----
ws.merge_cells("A1:C1")
ws["A1"] = "Case 02 — Debt Verification & Dispute Intake: Cost-Benefit Model"
ws["A1"].font = TITLE_FONT
ws.merge_cells("A2:C2")
ws["A2"] = "Source data: CFPB Consumer Complaint Database extract — Enhanced Recovery Company LLC, Debt collection product, Jan 1 – Mar 15 2015"
ws["A2"].font = Font(name=FONT, italic=True, size=9, color="52514E")
ws.merge_cells("A3:C3")
ws["A3"] = "Scoped to the single highest-volume company in the sample as a realistic company-level proxy. Swap in your own employer's ticket volume for a live business case."
ws["A3"].font = Font(name=FONT, italic=True, size=9, color="52514E")

r = 5
r = section(r, "1. OBSERVED SAMPLE DATA  (blue = pulled from the dataset)")
row_complaints_observed = r
r = label_val(r, "\"Debt is not mine\" complaints observed in sample window", 176, BLUE,
              comment="From analysis: sub-issue breakdown for the highest-volume company (Enhanced Recovery Company, LLC) in the #1 debt-collection issue. Source: consumer_complaints.csv (CFPB extract).")
row_window_days = r
r = label_val(r, "Sample window length (days)", 74, BLUE,
              comment="This company's complaints in the extract ranged 2015-01-01 to 2015-03-15.")
r += 1

r = section(r, "2. ANNUALIZED VOLUME  (formulas)")
row_annual_factor = r
r = label_val(r, "Annualization factor (365 / window days)", f"=365/B{row_window_days}", BLACK, "0.00")
row_annual_volume = r
r = label_val(r, "Estimated annualized \"not mine\" complaint volume", f"=B{row_complaints_observed}*B{row_annual_factor}", BLACK, "#,##0")
r += 1

r = section(r, "3. COST & IMPACT ASSUMPTIONS  (yellow = edit these for your own business case)")
row_cost_per = r
r = label_val(r, "Fully-loaded cost to process one complaint ($)", 150, BLUE, "$#,##0", fill=YELLOW,
              comment="Illustrative placeholder — covers agent research time, compliance documentation, and dispute handling. Replace with your company's actual complaint-handling cost per unit.")
row_reduction = r
r = label_val(r, "Expected reduction from adding an identity-verification gate", 0.40, BLUE, "0.0%", fill=YELLOW,
              comment="Team-estimated, conservative. Validate with a pilot on a sample of accounts before committing to this figure in a real business case.")
row_impl_cost = r
r = label_val(r, "One-time cost to build the verification step ($)", 45000, BLUE, "$#,##0", fill=YELLOW,
              comment="Illustrative placeholder for engineering + vendor data-match integration. Replace with a real estimate from engineering/procurement.")
row_run_cost = r
r = label_val(r, "Ongoing annual cost to run the verification step ($)", 12000, BLUE, "$#,##0", fill=YELLOW,
              comment="Illustrative placeholder — e.g. identity-verification API usage fees.")
r += 1

r = section(r, "4. PROJECTED IMPACT  (formulas — recalculates if you change section 3)")
row_avoided = r
r = label_val(r, "Complaints avoided per year", f"=B{row_annual_volume}*B{row_reduction}", BLACK, "#,##0")
row_gross_savings = r
r = label_val(r, "Gross annual savings", f"=B{row_avoided}*B{row_cost_per}", BLACK, "$#,##0")
row_net_savings = r
r = label_val(r, "Net annual savings (after run cost)", f"=B{row_gross_savings}-B{row_run_cost}", BLACK, "$#,##0")
row_payback = r
r = label_val(r, "Payback period (years)", f"=B{row_impl_cost}/B{row_net_savings}", BLACK, "0.00")
row_payback_months = r
r = label_val(r, "Payback period (months)", f"=B{row_payback}*12", BLACK, "0.0")
r += 1

r = section(r, "5. THREE-YEAR NET BENEFIT")
headers_row = r
for i, h in enumerate(["", "Year 1", "Year 2", "Year 3"]):
    c = ws.cell(row=headers_row, column=1+i if i == 0 else i+1, value=h)
ws.cell(row=headers_row, column=1, value="")
ws.cell(row=headers_row, column=2, value="Year 1").font = HEADER_FONT
ws.cell(row=headers_row, column=2).fill = HEADER_FILL
ws.cell(row=headers_row, column=3, value="Year 2 + 3 (combined)").font = HEADER_FONT
ws.cell(row=headers_row, column=3).fill = HEADER_FILL
r += 1
row_y1 = r
ws.cell(row=r, column=1, value="Net benefit").font = BLACK
ws.cell(row=r, column=2, value=f"=B{row_net_savings}-B{row_impl_cost}").number_format = "$#,##0"
ws.cell(row=r, column=3, value=f"=B{row_net_savings}*2").number_format = "$#,##0"
for col in (2, 3):
    ws.cell(row=r, column=col).border = BORDER
r += 1
row_3yr = r
ws.cell(row=r, column=1, value="Cumulative 3-year net benefit").font = BOLD
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
ws.cell(row=r, column=2, value=f"=B{row_y1}+C{row_y1}").number_format = "$#,##0"
ws.cell(row=r, column=2).font = BOLD
ws.cell(row=r, column=2).border = BORDER
r += 2

r = section(r, "LEGEND")
r = label_val(r, "Blue text", "Input pulled from the dataset or an editable assumption", BLACK)
r = label_val(r, "Yellow fill", "Key assumption — edit for your own business case", BLACK)
r = label_val(r, "Black text", "Formula — recalculates automatically", BLACK)

wb.save("case2_roi_model.xlsx")
print("Workbook saved.")
